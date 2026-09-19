"""Two-phase Excel import: validate (row-level error report) then commit.

§10: "Reject bad imports with a row-level error report; never silently
coerce." Every importer in this package follows the same shape:

1. `validate(rows)` -> (valid_rows, errors) — pure, no DB writes.
2. `commit(db, valid_rows, ...)` -> persists, wrapped in a transaction by
   the caller. `commit_valid_only=False` (the default) means the whole
   file is all-or-nothing: any error anywhere aborts the entire import.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Callable

import openpyxl
import pandas as pd


@dataclass
class RowError:
    row_number: int  # 1-based, matches the spreadsheet row (header = row 1)
    column: str
    value: Any
    error_code: str
    message: str


@dataclass
class ImportResult:
    total_rows: int
    valid_rows: list[dict]
    errors: list[RowError] = field(default_factory=list)

    @property
    def is_clean(self) -> bool:
        return len(self.errors) == 0


def read_workbook_rows(file_bytes: bytes, sheet_name: str | int = 0) -> list[dict]:
    """Reads one worksheet's rows as a list of {column: value} dicts.

    `sheet_name` is a *preferred* sheet when given as a string — the
    generated sample_masters.xlsx template has separate "staff"/"clients"
    sheets (plus a "Data Dictionary" one), so each endpoint asks for its
    own by name. A real firm's own export is typically a single sheet
    with an arbitrary name, though, so an unmatched string name falls
    back to the first sheet rather than erroring — an int index is
    always used as-is.
    """
    if isinstance(sheet_name, str):
        probe = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        try:
            if sheet_name not in probe.sheetnames:
                sheet_name = 0
        finally:
            probe.close()
    df = pd.read_excel(io.BytesIO(file_bytes), dtype=str, keep_default_na=False, sheet_name=sheet_name)
    df.columns = [str(c).strip() for c in df.columns]
    return df.to_dict(orient="records")


ColumnValidator = Callable[[int, dict], list[RowError]]


def run_validators(rows: list[dict], validators: list[ColumnValidator]) -> ImportResult:
    valid_rows: list[dict] = []
    errors: list[RowError] = []
    for idx, row in enumerate(rows, start=2):  # header is row 1
        row_errors: list[RowError] = []
        for validator in validators:
            row_errors.extend(validator(idx, row))
        if row_errors:
            errors.extend(row_errors)
        else:
            valid_rows.append(row)
    return ImportResult(total_rows=len(rows), valid_rows=valid_rows, errors=errors)


def required(field_name: str) -> ColumnValidator:
    def _v(row_number: int, row: dict) -> list[RowError]:
        value = row.get(field_name, "")
        if value is None or str(value).strip() == "":
            return [RowError(row_number, field_name, value, "REQUIRED", f"{field_name} is required")]
        return []

    return _v


def one_of(field_name: str, allowed: set[str], *, required_field: bool = True) -> ColumnValidator:
    def _v(row_number: int, row: dict) -> list[RowError]:
        value = str(row.get(field_name, "")).strip()
        if value == "":
            if required_field:
                return [RowError(row_number, field_name, value, "REQUIRED", f"{field_name} is required")]
            return []
        if value not in allowed:
            return [
                RowError(
                    row_number, field_name, value, "INVALID_ENUM",
                    f"{field_name}='{value}' is not one of {sorted(allowed)}",
                )
            ]
        return []

    return _v


def one_of_label(field_name: str, allowed_labels, *, required_field: bool = True) -> ColumnValidator:
    """Like `one_of`, but case-insensitive — for the friendly dropdown
    labels (e.g. "Senior Manager") a hand-typed or copy-pasted Excel cell
    may not match exactly on casing, unlike the raw enum-string columns
    `one_of` was built for.
    """
    lowered = {label.lower() for label in allowed_labels}

    def _v(row_number: int, row: dict) -> list[RowError]:
        value = str(row.get(field_name, "")).strip()
        if value == "":
            if required_field:
                return [RowError(row_number, field_name, value, "REQUIRED", f"{field_name} is required")]
            return []
        if value.lower() not in lowered:
            return [
                RowError(
                    row_number, field_name, value, "INVALID_ENUM",
                    f"{field_name}='{value}' is not one of {sorted(allowed_labels)}",
                )
            ]
        return []

    return _v


def numeric(field_name: str, *, required_field: bool = False) -> ColumnValidator:
    def _v(row_number: int, row: dict) -> list[RowError]:
        value = str(row.get(field_name, "")).strip()
        if value == "":
            if required_field:
                return [RowError(row_number, field_name, value, "REQUIRED", f"{field_name} is required")]
            return []
        try:
            float(value)
        except ValueError:
            return [RowError(row_number, field_name, value, "INVALID_NUMBER", f"{field_name}='{value}' is not numeric")]
        return []

    return _v


def build_error_workbook(errors: list[RowError]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Errors"
    ws.append(["row_number", "column", "value", "error_code", "message"])
    for e in errors:
        ws.append([e.row_number, e.column, str(e.value), e.error_code, e.message])
    ws.freeze_panes = "A2"
    for col_cells in ws.columns:
        max_len = max(len(str(c.value)) for c in col_cells if c.value is not None) if col_cells else 10
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
