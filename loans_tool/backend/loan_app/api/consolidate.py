"""Consolidate & Download — merge every ready upload of one report type
into a single dataset and download it as CSV, Parquet, or Excel.

Generalized from what used to be EAD-only ("EAD File Consolidation"):
every registered report type (EAD Files, Customer Master, Technical
Writeoff, Collection Report, Disbursement Report, ...) gets the same
page, parameterized by report_type. Consolidation itself is just
store.build_consolidated_df(), already generic (also used by EAD
Analytics and SQL Analytics for the other report types).
"""

from __future__ import annotations

import os
import tempfile
from datetime import UTC, datetime
from pathlib import Path

import polars as pl
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.templating import Jinja2Templates
from starlette.background import BackgroundTask

from fcmr_core.catalog import store
from fcmr_core.schemas.loader import available_report_types, label_for_report_type

router = APIRouter()
_templates_dir = Path(__file__).parent.parent / "web" / "templates"
templates = Jinja2Templates(directory=str(_templates_dir))


def _require_known_report_type(report_type: str) -> None:
    if report_type not in available_report_types():
        raise HTTPException(status_code=404, detail=f"Unknown report type: {report_type}")


def _consolidated_df(engagement_id: str | None, report_type: str) -> pl.DataFrame:
    return store.build_consolidated_df(engagement_id, report_type)


def _timestamped_filename(report_type: str, ext: str) -> str:
    ts = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    safe_label = label_for_report_type(report_type).replace(" ", "_")
    return f"{safe_label}_Consolidated_{ts}.{ext}"


@router.get("/dashboard/consolidate/{report_type}", response_class=HTMLResponse)
async def consolidate_page(request: Request, report_type: str):
    _require_known_report_type(report_type)
    engagement_id = request.session.get("engagement_id")
    uploads = store.list_uploads(engagement_id=engagement_id)
    ready = [u for u in uploads if u["report_type"] == report_type and u["status"] == "ready"]
    pending = [u for u in uploads if u["report_type"] == report_type and u["status"] == "mapping_pending"]
    return templates.TemplateResponse(
        request=request,
        name="consolidate.html",
        context={
            "report_type": report_type,
            "label": label_for_report_type(report_type),
            "ready": ready,
            "pending": pending,
            "total": len(ready),
        },
    )


@router.get("/dashboard/consolidate/{report_type}/download/csv")
async def consolidate_download_csv(request: Request, report_type: str):
    _require_known_report_type(report_type)
    engagement_id = request.session.get("engagement_id")
    df = _consolidated_df(engagement_id, report_type)
    if df.is_empty():
        raise HTTPException(status_code=404, detail="No ready files found to consolidate.")

    # Written straight to a temp file and served via FileResponse (streamed
    # by Starlette) rather than building the whole CSV as a Python string,
    # then again as bytes, then handing that one in-memory blob to
    # Response() as the entire body -- same fix as the SQL Analytics
    # export and EAD summary downloads, applied here for consistency (this
    # is exactly the "large data" case that matters most: a full
    # consolidated dataset, not a small aggregated summary).
    fd, tmp_name = tempfile.mkstemp(suffix=".csv")
    os.close(fd)
    tmp_path = Path(tmp_name)
    df.write_csv(tmp_path)

    return FileResponse(
        tmp_path,
        media_type="text/csv",
        filename=_timestamped_filename(report_type, "csv"),
        background=BackgroundTask(tmp_path.unlink, missing_ok=True),
    )


@router.get("/dashboard/consolidate/{report_type}/download/parquet")
async def consolidate_download_parquet(request: Request, report_type: str):
    _require_known_report_type(report_type)
    engagement_id = request.session.get("engagement_id")
    df = _consolidated_df(engagement_id, report_type)
    if df.is_empty():
        raise HTTPException(status_code=404, detail="No ready files found to consolidate.")

    fd, tmp_name = tempfile.mkstemp(suffix=".parquet")
    os.close(fd)
    tmp_path = Path(tmp_name)
    df.write_parquet(tmp_path)

    return FileResponse(
        tmp_path,
        media_type="application/octet-stream",
        filename=_timestamped_filename(report_type, "parquet"),
        background=BackgroundTask(tmp_path.unlink, missing_ok=True),
    )


@router.get("/dashboard/consolidate/{report_type}/download/excel")
async def consolidate_download_excel(request: Request, report_type: str):
    _require_known_report_type(report_type)
    engagement_id = request.session.get("engagement_id")
    df = _consolidated_df(engagement_id, report_type)
    if df.is_empty():
        raise HTTPException(status_code=404, detail="No ready files found to consolidate.")

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    label = label_for_report_type(report_type)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{label} Consolidated"[:31]  # Excel sheet-name length limit

    header_fill = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    data_font = Font(size=10)

    cols = df.columns
    for ci, col in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=ci, value=col.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = df.to_numpy(allow_copy=True)
    for ri, row in enumerate(rows, start=2):
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(
                row=ri, column=ci, value=None if (val is None or str(val) == "None") else val
            )
            cell.font = data_font

    for ci, col in enumerate(cols, start=1):
        max_len = max(len(col), 10)
        sample = df[col].drop_nulls().head(200).cast(pl.Utf8)
        if sample.len() > 0:
            max_len = max(max_len, sample.map_elements(len, return_dtype=pl.Int32).max() or 0)
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = f"{label} Consolidation Summary"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2["A3"] = "Total Rows"
    ws2["B3"] = len(df)
    ws2["A4"] = "Total Columns"
    ws2["B4"] = len(cols)
    ws2["A5"] = "Generated At"
    ws2["B5"] = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")
    if "_source_file" in df.columns:
        ws2["A7"] = "Source Files"
        ws2["A7"].font = Font(bold=True)
        source_files = df["_source_file"].unique().to_list()
        for i, fname in enumerate(source_files):
            ws2.cell(row=8 + i, column=1, value=fname)

    fd, tmp_name = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    tmp_path = Path(tmp_name)
    wb.save(tmp_path)

    return FileResponse(
        tmp_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=_timestamped_filename(report_type, "xlsx"),
        background=BackgroundTask(tmp_path.unlink, missing_ok=True),
    )
